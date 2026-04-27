import os
import sys
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from dotenv import load_dotenv
from huggingface_hub import HfApi, hf_hub_download, create_repo
from huggingface_hub.errors import RepositoryNotFoundError

load_dotenv()

# ─────────────────────────────────────────
# CONFIG
# ─────────────────────────────────────────
HF_TOKEN     = os.getenv("HF_TOKEN")
HF_REPO_ID   = "mariaamandadevina/jatim-curah-hujan-padi-2024"
HF_REPO_TYPE = "dataset"

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv("DATA_DIR", str(BASE_DIR / "__data__")))

CH_FILENAME   = "Data_Curah_Hujan.xlsx"
PADI_FILENAME = "Data_Padi.xlsx"
OUT_CSV       = DATA_DIR / "Data_Gabungan_Curah_Hujan_Padi.csv"
OUT_XLSX      = DATA_DIR / "Data_Gabungan_Curah_Hujan_Padi.xlsx"

BULAN_MAP = {
    "Januari": 1, "Februari": 2, "Maret": 3, "April": 4,
    "Mei": 5, "Juni": 6, "Juli": 7, "Agustus": 8,
    "September": 9, "Oktober": 10, "November": 11, "Desember": 12,
}


# ─────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────
def separator(title=""):
    if title:
        print(f"\n{'=' * 50}")
        print(f"  {title}")
        print(f"{'=' * 50}")
    else:
        print("=" * 50)


def ensure_repo_exists(api: HfApi):
    """Create the HuggingFace repo if it doesn't exist yet."""
    try:
        api.repo_info(repo_id=HF_REPO_ID, repo_type=HF_REPO_TYPE, token=HF_TOKEN)
        print(f"  ✓ Repo found: {HF_REPO_ID}")
    except RepositoryNotFoundError:
        print(f"  ℹ Repo not found. Creating: {HF_REPO_ID} ...")
        create_repo(
            repo_id=HF_REPO_ID,
            repo_type=HF_REPO_TYPE,
            private=False,
            token=HF_TOKEN,
        )
        print(f"  ✓ Repo created: https://huggingface.co/datasets/{HF_REPO_ID}")


def upload_files(api: HfApi, file_pairs: list[tuple[Path, str]]):
    """Upload a list of (local_path, repo_path) pairs to HuggingFace."""
    ensure_repo_exists(api)
    any_uploaded = False
    for local_path, repo_path in file_pairs:
        if not local_path.exists():
            print(f"  ⚠ Skipping (not found locally): {local_path.name}")
            continue
        print(f"  ↑ Uploading {local_path.name} → {repo_path} ...")
        api.upload_file(
            path_or_fileobj=str(local_path),
            path_in_repo=repo_path,
            repo_id=HF_REPO_ID,
            repo_type=HF_REPO_TYPE,
            token=HF_TOKEN,
        )
        print(f"  ✓ Done: {local_path.name}")
        any_uploaded = True
    if any_uploaded:
        print(f"\n  🔗 https://huggingface.co/datasets/{HF_REPO_ID}")


def parse_sheet_name(sheet_name):
    name = sheet_name.strip()
    if name.startswith("Kabupaten "):
        return name[len("Kabupaten "):].strip(), "A"
    elif name.startswith("Kota "):
        return name[len("Kota "):].strip(), "B"
    return name, "?"


def read_ch_sheet(df_raw):
    header_row = next(
        (i for i, v in enumerate(df_raw.iloc[:, 0]) if str(v).strip() == "Bulan (Month)"),
        None,
    )
    if header_row is None:
        return pd.DataFrame(columns=["Bulan", "Curah Hujan (mm)"])
    df = df_raw.iloc[header_row + 1:].copy()
    df.columns = ["Bulan", "Curah Hujan (mm)", *df_raw.columns[2:]]
    df = df[["Bulan", "Curah Hujan (mm)"]].dropna(subset=["Bulan"])
    df["Bulan"] = df["Bulan"].str.strip().map(BULAN_MAP)
    df = df.dropna(subset=["Bulan"])
    df["Bulan"] = df["Bulan"].astype(int)
    df["Curah Hujan (mm)"] = pd.to_numeric(df["Curah Hujan (mm)"], errors="coerce")
    return df.reset_index(drop=True)


def read_padi_sheet(df_raw):
    header_row = next(
        (i for i, v in enumerate(df_raw.iloc[:, 0]) if str(v).strip() == "Bulan (Month)"),
        None,
    )
    if header_row is None:
        return pd.DataFrame(columns=["Bulan", "Luas Panen (ha)", "Produktivitas (ton/ha)"])
    df = df_raw.iloc[header_row + 1:].copy()
    df.columns = ["Bulan", "Luas Panen (ha)", "Produksi GKG (ton)", "Produktivitas (ton/ha)", *df_raw.columns[4:]]
    df = df[["Bulan", "Luas Panen (ha)", "Produktivitas (ton/ha)"]].dropna(subset=["Bulan"])
    df["Bulan"] = df["Bulan"].str.strip().map(BULAN_MAP)
    df = df.dropna(subset=["Bulan"])
    df["Bulan"] = df["Bulan"].astype(int)
    df["Luas Panen (ha)"] = pd.to_numeric(df["Luas Panen (ha)"], errors="coerce")
    df["Produktivitas (ton/ha)"] = pd.to_numeric(df["Produktivitas (ton/ha)"], errors="coerce")
    return df.reset_index(drop=True)


def build_combined_df(ch_file: Path, padi_file: Path) -> pd.DataFrame:
    ch_sheets   = {k.strip(): v for k, v in pd.read_excel(ch_file,   sheet_name=None, header=None).items()}
    padi_sheets = {k.strip(): v for k, v in pd.read_excel(padi_file, sheet_name=None, header=None).items()}

    rows = []
    for sheet in sorted(ch_sheets.keys()):
        nama_area, jenis_area = parse_sheet_name(sheet)
        ch_df   = read_ch_sheet(ch_sheets[sheet])
        padi_df = read_padi_sheet(padi_sheets.get(sheet, pd.DataFrame()))
        merged  = pd.merge(ch_df, padi_df, on="Bulan", how="outer").sort_values("Bulan")
        merged.insert(0, "Nama Area",  nama_area)
        merged.insert(1, "Jenis Area", jenis_area)
        rows.append(merged)
        print(f"    ✓ {sheet} — {len(merged)} rows")

    combined = pd.concat(rows, ignore_index=True)
    combined = combined[["Bulan", "Nama Area", "Jenis Area",
                          "Curah Hujan (mm)", "Luas Panen (ha)", "Produktivitas (ton/ha)"]]
    return combined.sort_values(["Nama Area", "Bulan"]).reset_index(drop=True)


def save_xlsx(df: pd.DataFrame, path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Gabungan"  # type: ignore

    HEADER_FILL = PatternFill("solid", start_color="2E75B6", end_color="2E75B6")
    HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    EVEN_FILL   = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
    ODD_FILL    = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
    CENTER      = Alignment(horizontal="center", vertical="center")
    LEFT        = Alignment(horizontal="left",   vertical="center")
    THIN        = Side(style="thin", color="BFBFBF")
    BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    headers    = ["Bulan", "Nama Area", "Jenis Area (A=Kabupaten, B=Kota)",
                  "Curah Hujan (mm)", "Luas Panen (ha)", "Produktivitas (ton/ha)"]
    col_widths = [8, 18, 36, 20, 18, 22]

    for col_idx, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=h)  # type: ignore
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER
        cell.border    = BORDER
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w  # type: ignore
    ws.row_dimensions[1].height = 30  # type: ignore

    for r_idx, row in df.iterrows():
        fill   = EVEN_FILL if r_idx % 2 == 0 else ODD_FILL  # type: ignore
        values = [row["Bulan"], row["Nama Area"], row["Jenis Area"],
                  row["Curah Hujan (mm)"], row["Luas Panen (ha)"], row["Produktivitas (ton/ha)"]]
        aligns = [CENTER, LEFT, CENTER, CENTER, CENTER, CENTER]
        for c_idx, (val, aln) in enumerate(zip(values, aligns), 1):
            cell = ws.cell(row=r_idx + 2, column=c_idx, value=val)  # type: ignore
            cell.fill      = fill
            cell.font      = Font(name="Arial", size=10)
            cell.alignment = aln
            cell.border    = BORDER
            if c_idx == 1:
                cell.number_format = "0"
            elif c_idx in (4, 6):
                cell.number_format = "#,##0.00"
            elif c_idx == 5:
                cell.number_format = "#,##0"

    ws.freeze_panes = "A2"  # type: ignore
    wb.save(path)


# ─────────────────────────────────────────
# OPTION 1 — Upload local files to HuggingFace
# ─────────────────────────────────────────
def upload_local_files():
    separator("OPTION 1 — Upload local files to HuggingFace")

    # Collect every file in DATA_DIR the user might want to upload
    if not DATA_DIR.exists():
        print(f"  ✗ DATA_DIR not found: {DATA_DIR}")
        return

    all_files = [f for f in DATA_DIR.iterdir() if f.is_file()]
    if not all_files:
        print(f"  ✗ No files found in {DATA_DIR}")
        return

    print(f"\n  Files found in {DATA_DIR}:")
    for i, f in enumerate(all_files, 1):
        size_kb = f.stat().st_size / 1024
        print(f"    [{i}] {f.name}  ({size_kb:.1f} KB)")

    print("\n  Enter file numbers to upload (e.g. 1 2 3), or 'all':")
    choice = input("  > ").strip().lower()

    if choice == "all":
        selected = all_files
    else:
        try:
            indices = [int(x) - 1 for x in choice.split()]
            selected = [all_files[i] for i in indices]
        except (ValueError, IndexError):
            print("  ✗ Invalid selection.")
            return

    api = HfApi()
    pairs = [(f, f.name) for f in selected]
    upload_files(api, pairs)


# ─────────────────────────────────────────
# OPTION 2 — Download raw Excel files from HuggingFace
# ─────────────────────────────────────────
def download_raw_excels():
    separator("OPTION 2 — Download raw Excel files from HuggingFace")
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    for filename in [CH_FILENAME, PADI_FILENAME]:
        print(f"  ↓ Downloading {filename} ...")
        try:
            hf_hub_download(
                repo_id=HF_REPO_ID,
                filename=filename,
                repo_type=HF_REPO_TYPE,
                token=HF_TOKEN,
                local_dir=str(DATA_DIR),
            )
            print(f"  ✓ Saved: {DATA_DIR / filename}")
        except Exception as e:
            print(f"  ✗ Failed to download {filename}: {e}")


# ─────────────────────────────────────────
# OPTION 3 — Download combined CSV from HuggingFace
# ─────────────────────────────────────────
def download_combined_csv():
    separator("OPTION 3 — Download combined CSV from HuggingFace")
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    for filename in [OUT_CSV.name, OUT_XLSX.name]:
        print(f"  ↓ Downloading {filename} ...")
        try:
            hf_hub_download(
                repo_id=HF_REPO_ID,
                filename=filename,
                repo_type=HF_REPO_TYPE,
                token=HF_TOKEN,
                local_dir=str(DATA_DIR),
            )
            print(f"  ✓ Saved: {DATA_DIR / filename}")
        except Exception as e:
            print(f"  ✗ Failed to download {filename}: {e}")


# ─────────────────────────────────────────
# OPTION 4 — Combine downloaded raw Excel files
# ─────────────────────────────────────────
def combine_raw_excels():
    separator("OPTION 4 — Combine raw Excel files → CSV + XLSX")

    ch_file   = DATA_DIR / CH_FILENAME
    padi_file = DATA_DIR / PADI_FILENAME

    missing = []
    if not ch_file.exists():
        missing.append(ch_file.name)
    if not padi_file.exists():
        missing.append(padi_file.name)

    if missing:
        print(f"  ✗ Missing files: {', '.join(missing)}")
        print(f"    Run option 2 first to download them from HuggingFace.")
        return

    print(f"  Source (Curah Hujan): {ch_file}")
    print(f"  Source (Padi):        {padi_file}\n")

    df = build_combined_df(ch_file, padi_file)

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT_CSV, index=False)
    save_xlsx(df, OUT_XLSX)

    print(f"\n  ✓ Total rows : {len(df):,}")
    print(f"  ✓ CSV  saved : {OUT_CSV}")
    print(f"  ✓ XLSX saved : {OUT_XLSX}")

    print("\n  Summary by area:")
    print(
        df.groupby(["Nama Area", "Jenis Area"])
        .size()
        .reset_index(name="rows")
        .to_string(index=False)
    )


# ─────────────────────────────────────────
# OPTION 5 — Re-upload current local files (replace old on HF)
# ─────────────────────────────────────────
def update_files():
    separator("OPTION 5 — Update / re-upload local files to HuggingFace")

    # Re-use the same interactive upload logic
    upload_local_files()


# ─────────────────────────────────────────
# MENU
# ─────────────────────────────────────────
def main():
    separator()
    print("   JATIM CURAH HUJAN & PADI — DATA MANAGER")
    separator()
    print("  1. Upload local files to HuggingFace")
    print("  2. Download raw Excel files from HuggingFace")
    print("  3. Download combined CSV + XLSX from HuggingFace")
    print("  4. Combine downloaded raw Excel files → CSV + XLSX")
    print("  5. Update / re-upload local files to HuggingFace")
    print("  0. Exit")
    separator()

    choice = input("  Choose an option [0–5]: ").strip()

    actions = {
        "1": upload_local_files,
        "2": download_raw_excels,
        "3": download_combined_csv,
        "4": combine_raw_excels,
        "5": update_files,
    }

    if choice == "0":
        print("  Goodbye.")
        sys.exit(0)
    elif choice in actions:
        actions[choice]()
    else:
        print("  ✗ Invalid option. Please choose 0–5.")


if __name__ == "__main__":
    main()